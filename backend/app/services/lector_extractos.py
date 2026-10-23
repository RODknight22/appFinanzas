"""Lee extractos bancarios (CSV o Excel .xlsx y .xls) sin que el usuario configure nada.

Cada banco exporta a su manera: separador `;` o `,`, importes "1.234,56" o
"1,234.56", fechas "dd/mm/aaaa", líneas de cabecera con el titular antes de la
tabla, un solo importe con signo o dos columnas (cargo y abono)… Este módulo
lo detecta todo y devuelve siempre lo mismo: fecha, importe con signo y concepto.

No sabe nada de la base de datos: así se prueba con archivos de ejemplo y lo
podrá reutilizar cualquier otra fuente de movimientos.
"""

import csv
import io
import re
import zipfile
from collections import Counter
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.database import normalizar_texto
from app.services.errores import ReglaDeNegocioError

Celda = str | int | float | date | datetime | None

# Un extracto real cabe de sobra; el límite evita trabajar con un archivo que no lo es.
MAX_FILAS = 20_000
# La cabecera de la tabla suele estar en las primeras filas (antes va el titular, la cuenta…).
FILAS_BUSCAR_CABECERA = 40
LONGITUD_MAX_CONCEPTO = 500
# Un .xlsx es un ZIP: uno de pocos KB puede descomprimirse en gigas ("bomba ZIP") y
# colgar la app. Un extracto real ocupa pocos MB descomprimido y tiene una decena de piezas.
MAX_DESCOMPRIMIDO_EXCEL = 20_000_000
MAX_PIEZAS_EXCEL = 200

# Nombres de columna habituales en bancos de España y México, ya normalizados
# (minúsculas, sin tildes ni signos). Basta con que la cabecera empiece por uno de ellos.
SINONIMOS: dict[str, tuple[str, ...]] = {
    # Revolut da dos fechas: la de la compra ("started") va antes que la del cargo.
    "fecha": ("fecha", "f operacion", "f valor", "dia", "date", "started date", "completed date"),
    "concepto": (
        "concepto", "descripcion", "movimiento", "detalle", "description", "comercio",
        "establecimiento", "beneficiario", "observaciones", "referencia",
    ),
    "importe": ("importe", "monto", "cantidad", "amount", "valor"),
    "cargo": ("cargo", "cargos", "debe", "retiro", "retiros", "debito", "salida", "salidas"),
    "abono": ("abono", "abonos", "haber", "deposito", "depositos", "credito", "entrada", "entradas"),
    # Comisión aparte del importe (Revolut): también sale de la cuenta, así que se resta.
    "comision": ("comision", "comisiones", "fee", "fees"),
}
# Columnas numéricas que no son el movimiento: nunca se toman como importe.
NO_IMPORTE = ("saldo", "balance", "disponible")

_PATRON_FECHA = re.compile(r"^(\d{1,4})[/\-.](\d{1,2})[/\-.](\d{1,4})")
_PATRON_NUMERO = re.compile(r"^\d[\d.,]*$")


@dataclass(frozen=True)
class MovimientoLeido:
    fila: int  # número de fila en el archivo (empezando en 1), para los mensajes
    fecha: date
    importe: Decimal  # negativo = sale dinero de la cuenta
    concepto: str


@dataclass(frozen=True)
class ExtractoLeido:
    movimientos: list[MovimientoLeido]
    # Qué columna del archivo se usó para cada dato, p. ej. {"fecha": "F. Operación"}.
    columnas: dict[str, str]
    # Filas de la tabla que no parecían un movimiento (totales, líneas en blanco…).
    filas_ignoradas: int


# ---------------------------------------------------------------------------
# 1. Del archivo a una tabla de celdas
# ---------------------------------------------------------------------------
def _comprobar_zip(contenido: bytes) -> None:
    """Rechaza el Excel antes de abrirlo si descomprimido ocuparía demasiado.

    Basta con lo que declara el índice del ZIP: Python nunca extrae de una pieza
    más bytes de los que declara, así que mentir ahí solo produce un archivo dañado.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as archivo:
            piezas = archivo.infolist()
    except zipfile.BadZipFile as error:
        raise ReglaDeNegocioError("No se pudo abrir el Excel: el archivo parece dañado") from error
    if len(piezas) > MAX_PIEZAS_EXCEL or sum(p.file_size for p in piezas) > MAX_DESCOMPRIMIDO_EXCEL:
        raise ReglaDeNegocioError("El Excel es demasiado grande por dentro: exporta un periodo más corto o un CSV")


def _filas_excel(contenido: bytes) -> list[list[Celda]]:
    _comprobar_zip(contenido)
    # Importación local: openpyxl solo se carga si de verdad llega un Excel.
    from openpyxl import load_workbook

    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as error:  # openpyxl lanza tipos muy variados ante un archivo dañado
        raise ReglaDeNegocioError("No se pudo abrir el Excel: el archivo parece dañado") from error
    try:
        hoja = libro.active
        filas = []
        for fila in hoja.iter_rows(values_only=True):
            filas.append(list(fila))
            if len(filas) > MAX_FILAS:
                break
        return filas
    finally:
        libro.close()


def _filas_xls(contenido: bytes) -> list[list[Celda]]:
    """Lee el .xls antiguo (Excel 97-2003). No va comprimido, así que no hay bomba ZIP:
    lo que ocupa en memoria crece con el archivo, que ya llega limitado a 1 MB."""
    # Importación local, como openpyxl: xlrd solo se carga si llega un .xls.
    import xlrd

    try:
        # logfile: xlrd escribe sus avisos en la consola si no se le da otro sitio.
        libro = xlrd.open_workbook(file_contents=contenido, on_demand=True, logfile=io.StringIO())
    except Exception as error:  # xlrd también lanza tipos variados ante un archivo dañado
        raise ReglaDeNegocioError("No se pudo abrir el Excel: el archivo parece dañado") from error
    try:
        hoja = libro.sheet_by_index(0)

        def valor(celda: "xlrd.sheet.Cell") -> Celda:
            if celda.ctype == xlrd.XL_CELL_DATE:  # en .xls una fecha es un número con formato
                return xlrd.xldate_as_datetime(celda.value, libro.datemode)
            if celda.ctype in (xlrd.XL_CELL_TEXT, xlrd.XL_CELL_NUMBER):
                return celda.value
            return None  # vacías, booleanos y celdas con error

        return [[valor(c) for c in hoja.row(i)] for i in range(min(hoja.nrows, MAX_FILAS + 1))]
    finally:
        libro.release_resources()


def _decodificar(contenido: bytes) -> str:
    """UTF-8 si lo es; si no, Windows-1252 (lo que usan muchos bancos españoles)."""
    try:
        return contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        return contenido.decode("cp1252", errors="replace")


def _filas_csv(contenido: bytes) -> list[list[Celda]]:
    lineas = _decodificar(contenido).splitlines()[: MAX_FILAS + 1]

    def puntuar(separador: str) -> int:
        # El separador bueno parte la mayoría de filas en el mismo número de columnas (> 1).
        anchos = Counter(len(fila) for fila in csv.reader(lineas[:200], delimiter=separador))
        ancho, veces = max(anchos.items(), key=lambda par: (par[1], par[0]), default=(0, 0))
        return veces * ancho if ancho > 1 else 0

    separador = max(";,\t|", key=puntuar)
    return [list(fila) for fila in csv.reader(lineas, delimiter=separador)]


def _leer_tabla(contenido: bytes) -> list[list[Celda]]:
    if not contenido.strip():
        raise ReglaDeNegocioError("El archivo está vacío")
    if contenido.startswith(b"PK"):  # .xlsx es un ZIP
        return _filas_excel(contenido)
    if contenido.startswith(b"\xd0\xcf\x11\xe0"):  # .xls antiguo (formato OLE), p. ej. Santander
        return _filas_xls(contenido)
    if contenido.lstrip()[:1] == b"<":  # algunos bancos llaman ".xls" a una página HTML
        raise ReglaDeNegocioError("El archivo es una página web, no una hoja de cálculo: expórtalo como CSV o .xlsx")
    return _filas_csv(contenido)


# ---------------------------------------------------------------------------
# 2. Interpretar celdas sueltas
# ---------------------------------------------------------------------------
def _texto(celda: Celda) -> str:
    return "" if celda is None else str(celda).strip()


def _partes_fecha(celda: Celda) -> tuple[int, int, int] | None:
    match = _PATRON_FECHA.match(_texto(celda))
    return tuple(int(parte) for parte in match.groups()) if match else None  # type: ignore[return-value]


def _dia_primero(celdas: Iterable[Celda]) -> bool:
    """¿"03/04" es 3 de abril (España, México) o 4 de marzo (EE. UU.)?

    Lo decide cualquier fecha cuyo primer número pase de 12; si ninguna lo
    aclara, se asume día primero, lo habitual en los bancos de Rodrigo.
    """
    for celda in celdas:
        partes = _partes_fecha(celda)
        if partes is None or partes[0] > 31:  # aaaa-mm-dd: no aporta nada
            continue
        if partes[0] > 12:
            return True
        if partes[1] > 12:
            return False
    return True


def _a_fecha(celda: Celda, dia_primero: bool) -> date | None:
    if isinstance(celda, datetime):
        return celda.date()
    if isinstance(celda, date):
        return celda
    partes = _partes_fecha(celda)
    if partes is None:
        return None
    a, b, c = partes
    if a > 31:  # aaaa-mm-dd
        anio, mes, dia = a, b, c
    else:
        dia, mes = (a, b) if dia_primero else (b, a)
        anio = c + 2000 if c < 100 else c
    try:
        return date(anio, mes, dia)
    except ValueError:
        return None


def _a_importe(celda: Celda) -> Decimal | None:
    """Convierte "1.234,56", "1,234.56", "-12,30 €", "(5.00)" o "12,30-" en Decimal."""
    if isinstance(celda, bool) or celda is None:
        return None
    if isinstance(celda, int | float):
        return Decimal(str(celda)).quantize(Decimal("0.01"))
    texto = re.sub(r"[\s €$]|eur|mxn|usd", "", _texto(celda).lower()).replace("−", "-")
    negativo = texto.startswith("-") or texto.endswith("-") or (texto.startswith("(") and texto.endswith(")"))
    texto = texto.strip("+-()")
    if not _PATRON_NUMERO.match(texto):
        return None
    if "." in texto and "," in texto:
        # El que aparece último es el decimal: "1.234,56" o "1,234.56".
        decimal = "," if texto.rfind(",") > texto.rfind(".") else "."
    elif texto.count(",") + texto.count(".") == 1 and not re.search(r"[.,]\d{3}$", texto):
        # Un solo separador sin tres cifras detrás ("12,5", "7.99"): es el decimal.
        decimal = "," if "," in texto else "."
    else:
        # "1.234" o "1,234,567": solo separadores de miles (nadie tiene 3 decimales en euros).
        decimal = None
    miles = {",": ".", ".": ",", None: ".,"}[decimal]
    for separador in miles:
        texto = texto.replace(separador, "")
    if decimal == ",":
        texto = texto.replace(",", ".")
    try:
        valor = Decimal(texto).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None
    return -valor if negativo else valor


# ---------------------------------------------------------------------------
# 3. Encontrar la tabla y el papel de cada columna
# ---------------------------------------------------------------------------
def _normalizar_cabecera(celda: Celda) -> str:
    texto = normalizar_texto(_texto(celda)) or ""
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texto).split())


def _rol(cabecera: str) -> tuple[str, int] | None:
    """El dato que contiene una columna y lo típico que es su nombre (0 = el más típico)."""
    if not cabecera or any(palabra in cabecera for palabra in NO_IMPORTE):
        return None
    for rol, sinonimos in SINONIMOS.items():
        for posicion, sinonimo in enumerate(sinonimos):
            if cabecera == sinonimo or cabecera.startswith(sinonimo + " "):
                return rol, posicion
    return None


def _roles_por_cabecera(fila: list[Celda]) -> dict[str, int]:
    """Columna de cada dato según los nombres de la cabecera.

    Si varias encajan gana el nombre más típico ("Concepto" antes que "Referencia")
    y, entre fechas, la de operación antes que la "fecha valor".
    """
    candidatas: dict[str, list[tuple[int, bool, int]]] = {}
    for indice, celda in enumerate(fila):
        cabecera = _normalizar_cabecera(celda)
        encontrado = _rol(cabecera)
        if encontrado is not None:
            rol, posicion = encontrado
            candidatas.setdefault(rol, []).append((posicion, "valor" in cabecera, indice))
    return {rol: min(opciones)[2] for rol, opciones in candidatas.items()}


def _es_cabecera(roles: dict[str, int]) -> bool:
    return "fecha" in roles and ("importe" in roles or "cargo" in roles or "abono" in roles)


def _proporcion(valores: list[Celda], prueba: Callable[[Celda], object]) -> float:
    """Qué parte de las celdas no vacías entiende `prueba` (devuelve algo distinto de None)."""
    llenos = [v for v in valores if _texto(v)]
    return sum(prueba(v) is not None for v in llenos) / len(llenos) if llenos else 0


def _roles_por_contenido(datos: list[list[Celda]], roles: dict[str, int]) -> dict[str, int]:
    """Completa los roles que la cabecera no aclaró mirando qué contiene cada columna."""
    ancho = max((len(fila) for fila in datos), default=0)
    columnas = [[fila[i] if i < len(fila) else None for fila in datos] for i in range(ancho)]
    usadas = set(roles.values())
    libres = [i for i in range(ancho) if i not in usadas]

    if "fecha" not in roles:
        for i in libres:
            if _proporcion(columnas[i], lambda v: _a_fecha(v, True)) > 0.8:
                roles["fecha"] = i
                break
    if not {"importe", "cargo", "abono"} & roles.keys():
        # La primera columna numérica: el saldo, si lo hay, suele ir después del importe.
        for i in libres:
            if i not in roles.values() and _proporcion(columnas[i], _a_importe) > 0.8:
                roles["importe"] = i
                break
    if "concepto" not in roles:
        # El texto más variado de las columnas restantes (el concepto casi nunca se repite igual).
        candidatas = [
            i for i in libres
            if i not in roles.values() and _proporcion(columnas[i], _a_importe) < 0.5
        ]
        if candidatas:
            roles["concepto"] = max(candidatas, key=lambda i: len({_texto(v) for v in columnas[i]}))
    return roles


def _localizar_tabla(filas: list[list[Celda]]) -> tuple[int, dict[str, int]]:
    """Devuelve dónde empiezan los datos y qué columna es cada dato.

    Primero busca una fila de cabecera reconocible; si el archivo no tiene
    cabecera, los datos empiezan en la primera fila que lleve una fecha.
    """
    for indice, fila in enumerate(filas[:FILAS_BUSCAR_CABECERA]):
        roles = _roles_por_cabecera(fila)
        if _es_cabecera(roles):
            return indice + 1, _roles_por_contenido(filas[indice + 1:], roles)
    for indice, fila in enumerate(filas[:FILAS_BUSCAR_CABECERA]):
        if any(_a_fecha(celda, True) for celda in fila):
            return indice, _roles_por_contenido(filas[indice:], {})
    return 0, {}


# ---------------------------------------------------------------------------
# 4. Punto de entrada
# ---------------------------------------------------------------------------
def leer_extracto(contenido: bytes) -> ExtractoLeido:
    """Lee un CSV, .xlsx o .xls de cualquier banco y devuelve sus movimientos."""
    filas = _leer_tabla(contenido)
    inicio, roles = _localizar_tabla(filas)
    if "fecha" not in roles or not {"importe", "cargo", "abono"} & roles.keys():
        raise ReglaDeNegocioError(
            "No se encontraron las columnas de fecha e importe: ¿es un extracto de movimientos?"
        )

    def celda(fila: list[Celda], rol: str) -> Celda:
        indice = roles.get(rol)
        return fila[indice] if indice is not None and indice < len(fila) else None

    datos = filas[inicio:]
    dia_primero = _dia_primero(celda(fila, "fecha") for fila in datos)
    movimientos: list[MovimientoLeido] = []
    ignoradas = 0
    for numero, fila in enumerate(datos, start=inicio + 1):
        fecha = _a_fecha(celda(fila, "fecha"), dia_primero)
        if "importe" in roles:
            importe = _a_importe(celda(fila, "importe"))
        else:
            # Cargo y abono por separado: unos bancos los ponen en positivo y otros el cargo en negativo.
            cargo, abono = _a_importe(celda(fila, "cargo")), _a_importe(celda(fila, "abono"))
            importe = None if cargo is None and abono is None else abs(abono or 0) - abs(cargo or 0)
        comision = _a_importe(celda(fila, "comision"))
        if importe is not None and comision:
            importe -= abs(comision)  # así entra la cuota del plan, que llega con importe 0
        if fecha is None or not importe:  # totales, notas al pie, líneas vacías o importe 0
            if any(_texto(c) for c in fila):
                ignoradas += 1
            continue
        concepto = " ".join(_texto(celda(fila, "concepto")).split())[:LONGITUD_MAX_CONCEPTO]
        movimientos.append(MovimientoLeido(numero, fecha, importe, concepto))

    if not movimientos:
        raise ReglaDeNegocioError("El archivo no contiene ningún movimiento reconocible")
    cabecera = filas[inicio - 1] if inicio > 0 and _es_cabecera(_roles_por_cabecera(filas[inicio - 1])) else None
    columnas = {
        rol: (_texto(cabecera[i]) if cabecera and i < len(cabecera) else f"Columna {i + 1}")
        for rol, i in roles.items()
    }
    return ExtractoLeido(movimientos, columnas, ignoradas)
